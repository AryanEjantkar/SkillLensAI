import logging
import ipaddress
import socket
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit

import httpx
from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field, field_validator
from pypdf import PdfReader
from starlette.staticfiles import StaticFiles

from backend.src.config.settings import MAX_PROMPT_CHARS, MODEL_ID
from backend.src.services.gemma_service import (
    GemmaConfigurationError,
    GemmaRequestError,
    GemmaService,
)

logger = logging.getLogger(__name__)
app = FastAPI(title="Gemma 4 Prompt Test", version="0.1.0")
gemma_service = GemmaService()
STATIC_DIRECTORY = Path(__file__).parent / "static"
INDEX_FILE = STATIC_DIRECTORY / "index.html"
FRONTEND_DIRECTORY = STATIC_DIRECTORY / "frontend"
FRONTEND_INDEX_FILE = FRONTEND_DIRECTORY / "index.html"
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_TEXT_CHARS = 40_000
PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"
PROMPT_DIRECTIONS = {
    "Explain": "Explain the material clearly. Start with the core idea, then explain important details and one useful example.",
    "Flashcards": "Create concise study flashcards for this topic. Format each card strictly as:\nQ: [Clear question or concept]\nA: [Clear, comprehensive answer]",
    "Quiz": "Create a multiple-choice quiz with 4 to 5 questions for this topic. Format each question strictly as:\n1. [Question text]\nA) [Option A]\nB) [Option B]\nC) [Option C]\nD) [Option D]\nAnswer: [Correct Option Letter]\nExplanation: [Brief explanation]",
    "All": "Create a complete study set with three distinct sections: 1) Explanation, 2) Flashcards (formatted strictly as Q: and A:), and 3) Quiz (formatted with questions 1-4, options A-D, Answer: [Letter], and Explanation:).",
}

if FRONTEND_DIRECTORY.is_dir():
    app.mount(
        "/frontend",
        StaticFiles(directory=FRONTEND_DIRECTORY),
        name="frontend-assets",
    )


class PromptRequest(BaseModel):
    prompt: str = Field(min_length=1, max_length=MAX_PROMPT_CHARS)

    @field_validator("prompt")
    @classmethod
    def prompt_must_not_be_blank(cls, value: str) -> str:
        value = value.strip()
        if not value:
            raise ValueError("Prompt cannot be blank")
        return value


class PromptResponse(BaseModel):
    model: str
    response: str


class LinkRequest(BaseModel):
    url: str = Field(min_length=8, max_length=2048)
    prompt: str = Field(default="", max_length=MAX_PROMPT_CHARS)
    study_mode: str = Field(default="Explain", pattern="^(Explain|Flashcards|Quiz|All)$")

    @field_validator("url")
    @classmethod
    def url_must_be_http(cls, value: str) -> str:
        parsed = urlsplit(value.strip())
        if parsed.scheme not in {"http", "https"} or not parsed.hostname:
            raise ValueError("Enter a valid HTTP or HTTPS link")
        if parsed.username or parsed.password:
            raise ValueError("Links with embedded credentials are not accepted")
        return value.strip()


class VisibleTextParser(HTMLParser):
    _ignored_tags = {"head", "script", "style", "noscript", "svg"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._ignored_depth = 0
        self._chunks: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._ignored_tags:
            self._ignored_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in self._ignored_tags and self._ignored_depth:
            self._ignored_depth -= 1

    def handle_data(self, data: str) -> None:
        if not self._ignored_depth and data.strip():
            self._chunks.append(data.strip())

    def text(self) -> str:
        return "\n".join(self._chunks)


def _check_context_size(text: str) -> str:
    text = text.strip()
    if not text:
        raise HTTPException(status_code=400, detail="No readable text was found in that source.")
    if len(text) > MAX_EXTRACTED_TEXT_CHARS:
        raise HTTPException(
            status_code=413,
            detail=f"Extracted source text is limited to {MAX_EXTRACTED_TEXT_CHARS:,} characters.",
        )
    return text


def _extract_document_text(filename: str, data: bytes) -> str:
    extension = Path(filename).suffix.lower()
    try:
        if extension == ".txt":
            text = data.decode("utf-8-sig")
        elif extension == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise HTTPException(status_code=400, detail="The uploaded file is not a valid PDF.")
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise HTTPException(status_code=400, detail="Password-protected PDFs are not supported.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif extension == ".docx":
            if not data.startswith(b"PK"):
                raise HTTPException(status_code=400, detail="The uploaded file is not a valid DOCX document.")
            document = Document(BytesIO(data))
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            tables = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            text = "\n".join(paragraphs + tables)
        else:
            raise HTTPException(
                status_code=415,
                detail="Supported documents are PDF, DOCX, and TXT files.",
            )
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(status_code=400, detail="The document could not be read.") from error
    return _check_context_size(text)


def _validate_public_url(url: str) -> tuple[str, int]:
    parsed = urlsplit(url)
    hostname = parsed.hostname or ""
    try:
        port = parsed.port or (443 if parsed.scheme == "https" else 80)
    except ValueError as error:
        raise HTTPException(status_code=400, detail="The link has an invalid port.") from error
    expected_port = 443 if parsed.scheme == "https" else 80
    if port != expected_port:
        raise HTTPException(status_code=400, detail="Links may only use the standard HTTP or HTTPS port.")
    if hostname.lower() == "localhost" or hostname.lower().endswith((".localhost", ".local")):
        raise HTTPException(status_code=400, detail="Links to local network hosts are not accepted.")

    try:
        addresses = [ipaddress.ip_address(hostname)]
    except ValueError:
        try:
            addresses = [
                ipaddress.ip_address(result[4][0])
                for result in socket.getaddrinfo(hostname, port, type=socket.SOCK_STREAM)
            ]
        except (OSError, ValueError) as error:
            raise HTTPException(status_code=400, detail="The link host could not be resolved.") from error
    if not addresses or any(not address.is_global for address in addresses):
        raise HTTPException(status_code=400, detail="Links to private or reserved network addresses are blocked.")
    return hostname, port


def _fetch_link_text(url: str) -> str:
    _validate_public_url(url)
    content = bytearray()
    try:
        with httpx.Client(
            timeout=httpx.Timeout(12.0, connect=4.0),
            follow_redirects=False,
            trust_env=False,
            headers={"User-Agent": "SkillLensStudy/1.0"},
        ) as client:
            with client.stream("GET", url) as response:
                if 300 <= response.status_code < 400:
                    raise HTTPException(status_code=400, detail="Redirecting links are not supported.")
                if response.status_code >= 400:
                    raise HTTPException(status_code=422, detail="The linked page could not be fetched.")
                content_type = response.headers.get("content-type", "").split(";", 1)[0].lower()
                for chunk in response.iter_bytes():
                    content.extend(chunk)
                    if len(content) > MAX_UPLOAD_BYTES:
                        raise HTTPException(status_code=413, detail="Linked content exceeds the 10 MB limit.")
    except HTTPException:
        raise
    except httpx.TimeoutException as error:
        raise HTTPException(status_code=504, detail="Fetching the link timed out.") from error
    except httpx.HTTPError as error:
        raise HTTPException(status_code=422, detail="The linked page could not be fetched.") from error

    data = bytes(content)
    if content_type == "application/pdf":
        return _extract_document_text("linked.pdf", data)
    if content_type == "text/plain":
        try:
            return _check_context_size(data.decode("utf-8-sig"))
        except UnicodeDecodeError as error:
            raise HTTPException(status_code=415, detail="The linked text is not UTF-8 encoded.") from error
    if content_type not in {"text/html", "application/xhtml+xml"}:
        raise HTTPException(status_code=415, detail="Links must point to an HTML, plain-text, or PDF page.")

    parser = VisibleTextParser()
    parser.feed(data.decode("utf-8", errors="replace"))
    return _check_context_size(parser.text())


def _study_prompt(prompt: str, study_mode: str, source: str, source_text: str | None = None) -> str:
    default_requests = {
        "image": "Summarize this image and describe its important details.",
        "document": "Summarize this document and explain its key points.",
        "link": "Summarize this page and explain its key points.",
    }
    question = prompt.strip() or default_requests[source]
    instruction = PROMPT_DIRECTIONS.get(study_mode, PROMPT_DIRECTIONS["Explain"])
    parts = [instruction, f"User request: {question}"]
    if source_text:
        parts.append(f"Extracted {source} content:\n{source_text}")
    elif source == "image":
        parts.append("Analyze the attached image as the source material.")
    return "\n\n".join(parts)


def _generate_with_gemma(
    prompt: str,
    image_data: bytes | None = None,
    image_mime_type: str | None = None,
) -> str:
    try:
        if image_data is None:
            return gemma_service.generate(prompt)
        return gemma_service.generate(prompt, image_data, image_mime_type)
    except GemmaConfigurationError as error:
        raise HTTPException(
            status_code=503,
            detail="Server configuration is missing GEMINI_API_KEY.",
        ) from error
    except GemmaRequestError as error:
        logger.warning("Gemma request failed for model %s", MODEL_ID)
        raise HTTPException(
            status_code=502,
            detail=f"Gemma could not process this request with {MODEL_ID}. Check the input and try again.",
        ) from error


async def _read_upload(file: UploadFile) -> bytes:
    data = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Uploads are limited to 10 MB.")
    if not data:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    return data


@app.get("/", include_in_schema=False)
def index() -> FileResponse:
    if FRONTEND_INDEX_FILE.is_file():
        return FileResponse(FRONTEND_INDEX_FILE)
    return FileResponse(INDEX_FILE)


@app.post("/api/generate", response_model=PromptResponse)
def generate(request: PromptRequest) -> PromptResponse:
    return PromptResponse(model=MODEL_ID, response=_generate_with_gemma(request.prompt))


@app.post("/api/generate/image", response_model=PromptResponse)
async def generate_from_image(
    file: UploadFile = File(...),
    prompt: str = Form(default="", max_length=MAX_PROMPT_CHARS),
    study_mode: str = Form(default="Explain", pattern="^(Explain|Flashcards|Quiz|All)$"),
) -> PromptResponse:
    filename = Path(file.filename or "").name
    if Path(filename).suffix.lower() != ".png":
        raise HTTPException(status_code=415, detail="PNG images are currently supported.")
    if file.content_type not in {"image/png", "application/octet-stream"}:
        raise HTTPException(status_code=415, detail="Upload a PNG image.")
    image_data = await _read_upload(file)
    if not image_data.startswith(PNG_SIGNATURE):
        raise HTTPException(status_code=400, detail="The uploaded file is not a valid PNG image.")

    gemma_prompt = _study_prompt(prompt, study_mode, "image")
    response = _generate_with_gemma(gemma_prompt, image_data, "image/png")
    return PromptResponse(model=MODEL_ID, response=response)


@app.post("/api/generate/document", response_model=PromptResponse)
async def generate_from_document(
    file: UploadFile = File(...),
    prompt: str = Form(default="", max_length=MAX_PROMPT_CHARS),
    study_mode: str = Form(default="Explain", pattern="^(Explain|Flashcards|Quiz|All)$"),
) -> PromptResponse:
    filename = Path(file.filename or "document").name
    document_data = await _read_upload(file)
    document_text = _extract_document_text(filename, document_data)
    gemma_prompt = _study_prompt(prompt, study_mode, "document", document_text)
    response = _generate_with_gemma(gemma_prompt)
    return PromptResponse(model=MODEL_ID, response=response)


@app.post("/api/generate/link", response_model=PromptResponse)
def generate_from_link(request: LinkRequest) -> PromptResponse:
    page_text = _fetch_link_text(request.url)
    gemma_prompt = _study_prompt(request.prompt, request.study_mode, "link", page_text)
    response = _generate_with_gemma(gemma_prompt)
    return PromptResponse(model=MODEL_ID, response=response)