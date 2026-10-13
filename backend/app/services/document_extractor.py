"""Utilities for extracting text from files, PDFs, DOCX, and URLs safely."""

import ipaddress
import socket
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit

import httpx
from docx import Document
from fastapi import HTTPException
from pypdf import PdfReader

from backend.app.core.config import MAX_EXTRACTED_TEXT_CHARS, MAX_UPLOAD_BYTES

PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"


class VisibleTextParser(HTMLParser):
    _ignored_tags = {"head", "script", "style", "noscript", "svg"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._ignored_depth = 0
        self._chunks: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._ignored_tags:
            self._ignored_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in self._ignored_tags and self._ignored_depth:
            self._ignored_depth -= 1

    def handle_data(self, data: str) -> None:
        if not self._ignored_depth and data.strip():
            self._chunks.append(data.strip())

    def text(self) -> str:
        return "\n".join(self._chunks)


def check_context_size(text: str) -> str:
    text = text.strip()
    if not text:
        raise HTTPException(status_code=400, detail="No readable text was found in that source.")
    if len(text) > MAX_EXTRACTED_TEXT_CHARS:
        raise HTTPException(
            status_code=413,
            detail=f"Extracted source text is limited to {MAX_EXTRACTED_TEXT_CHARS:,} characters.",
        )
    return text


def extract_document_text(filename: str, data: bytes) -> str:
    extension = Path(filename).suffix.lower()
    try:
        if extension == ".txt":
            text = data.decode("utf-8-sig")
        elif extension == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise HTTPException(status_code=400, detail="The uploaded file is not a valid PDF.")
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise HTTPException(status_code=400, detail="Password-protected PDFs are not supported.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif extension == ".docx":
            if not data.startswith(b"PK"):
                raise HTTPException(status_code=400, detail="The uploaded file is not a valid DOCX document.")
            document = Document(BytesIO(data))
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            tables = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            text = "\n".join(paragraphs + tables)
        else:
            raise HTTPException(
                status_code=415,
                detail="Supported documents are PDF, DOCX, and TXT files.",
            )
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(status_code=400, detail="The document could not be read.") from error
    return check_context_size(text)


def validate_public_url(url: str) -> tuple[str, int]:
    parsed = urlsplit(url)
    hostname = parsed.hostname or ""
    try:
        port = parsed.port or (443 if parsed.scheme == "https" else 80)
    except ValueError as error:
        raise HTTPException(status_code=400, detail="The link has an invalid port.") from error
    expected_port = 443 if parsed.scheme == "https" else 80
    if port != expected_port:
        raise HTTPException(status_code=400, detail="Links may only use the standard HTTP or HTTPS port.")
    if hostname.lower() == "localhost" or hostname.lower().endswith((".localhost", ".local")):
        raise HTTPException(status_code=400, detail="Links to local network hosts are not accepted.")

    try:
        addresses = [ipaddress.ip_address(hostname)]
    except ValueError:
        try:
            addresses = [
                ipaddress.ip_address(result[4][0])
                for result in socket.getaddrinfo(hostname, port, type=socket.SOCK_STREAM)
            ]
        except (OSError, ValueError) as error:
            raise HTTPException(status_code=400, detail="The link host could not be resolved.") from error
    if not addresses or any(not address.is_global for address in addresses):
        raise HTTPException(status_code=400, detail="Links to private or reserved network addresses are blocked.")
    return hostname, port


def fetch_link_text(url: str) -> str:
    validate_public_url(url)
    content = bytearray()
    try:
        with httpx.Client(
            timeout=httpx.Timeout(12.0, connect=4.0),
            follow_redirects=False,
            trust_env=False,
            headers={"User-Agent": "SkillLensStudy/1.0"},
        ) as client:
            with client.stream("GET", url) as response:
                if 300 <= response.status_code < 400:
                    raise HTTPException(status_code=400, detail="Redirecting links are not supported.")
                if response.status_code >= 400:
                    raise HTTPException(status_code=422, detail="The linked page could not be fetched.")
                content_type = response.headers.get("content-type", "").split(";", 1)[0].lower()
                for chunk in response.iter_bytes():
                    content.extend(chunk)
                    if len(content) > MAX_UPLOAD_BYTES:
                        raise HTTPException(status_code=413, detail="Linked content exceeds the 10 MB limit.")
    except HTTPException:
        raise
    except httpx.TimeoutException as error:
        raise HTTPException(status_code=504, detail="Fetching the link timed out.") from error
    except httpx.HTTPError as error:
        raise HTTPException(status_code=422, detail="The linked page could not be fetched.") from error

    data = bytes(content)
    if content_type == "application/pdf":
        return extract_document_text("linked.pdf", data)
    if content_type == "text/plain":
        try:
            return check_context_size(data.decode("utf-8-sig"))
        except UnicodeDecodeError as error:
            raise HTTPException(status_code=415, detail="The linked text is not UTF-8 encoded.") from error
    if content_type not in {"text/html", "application/xhtml+xml"}:
        raise HTTPException(status_code=415, detail="Links must point to an HTML, plain-text, or PDF page.")

    parser = VisibleTextParser()
    parser.feed(data.decode("utf-8", errors="replace"))
    return check_context_size(parser.text())
